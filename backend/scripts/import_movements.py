"""
Import movements from Excel file and create linked Document records.

Usage:
    docker compose exec backend python scripts/import_movements.py <file.xlsx>
    docker compose exec backend python scripts/import_movements.py <file.xlsx> --check

--check prints first 5 rows with column values so you can verify the mapping.

Excel format (header in row 2, data from row 3):
    A (0)  - Дата внесення       → entry_date
    B (1)  - Найменування        → item_name
    C (2)  - не працює           → SKIP
    D (3)  - Одиниця виміру      → unit_of_measure
    E (4)  - Категорія           → category
    F (5)  - Надійшло            → qty_in
    G (6)  - Вибуло              → qty_out
    H (7)  - Звідки              → from_unit
    I (8)  - Куди                → to_unit
    J (9)  - МВО звідки          → mvo_from_id
    K (10) - МВО куди            → mvo_to_id
    L (11) - Підстава            → basis
    M (12) - Дата документа      → doc_date
    N (13) - № документа         → doc_number
    O (14) - № без префіксу      → SKIP
    P (15) - Примітка            → serial_number
    Q (16) - Виконавець          → SKIP
    R (17) - Код номенклатури    → nomenclature_code
    S (18) - Термін дії накладної→ SKIP
    T (19) - Ціна                → price
    U (20) - Служба              → service
    V (21) - Поле 10             → SKIP
    W (22) - Поле 11             → SKIP
    X (23) - Поле 12 (№ картки)  → item_card_num (FK check)
    Y (24) - Тип документа       → doc_type (Excel label, stored as-is)
    Z (25) - Тип операції        → mapped to our doc_type for Document
    AA(26) - Категорія отримувача→ recipient_category
"""

import sys
import os
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

import openpyxl
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

sys.path.insert(0, '/app')
from app.models import Movement, Person, Document

DATABASE_URL = os.environ['DATABASE_URL']
engine = create_engine(DATABASE_URL)

OP_TYPE_MAP = {
    'надходження':           'надходження',
    'внутрішнє переміщення': 'переміщення',
    'переміщення':           'переміщення',
}

COL_LABELS = {
    0:  'A  entry_date',
    1:  'B  item_name',
    2:  'C  (не працює)',
    3:  'D  unit_of_measure',
    4:  'E  category',
    5:  'F  qty_in',
    6:  'G  qty_out',
    7:  'H  from_unit',
    8:  'I  to_unit',
    9:  'J  mvo_from',
    10: 'K  mvo_to',
    11: 'L  basis',
    12: 'M  doc_date',
    13: 'N  doc_number',
    14: 'O  (№ без префіксу)',
    15: 'P  serial_number',
    16: 'Q  (виконавець)',
    17: 'R  nomenclature_code',
    18: 'S  (термін дії)',
    19: 'T  price',
    20: 'U  service',
    21: 'V  (Поле 10)',
    22: 'W  (Поле 11)',
    23: 'X  item_card_num',
    24: 'Y  doc_type_excel',
    25: 'Z  op_type',
    26: 'AA recipient_category',
}


def clean(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def parse_date(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if not s:
        return None
    for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return s


def parse_decimal(val):
    if val is None:
        return None
    try:
        return Decimal(str(val).strip().replace(',', '.'))
    except (InvalidOperation, ValueError):
        return None


def col(row, idx):
    return row[idx].value if idx < len(row) else None


def build_person_map(session):
    mapping = {}
    for p in session.query(Person).all():
        if p.first_name and p.last_name:
            mapping[f"{p.first_name} {p.last_name}".strip().lower()] = p.id
        if p.search_name:
            mapping[p.search_name.strip().lower()] = p.id
    return mapping


def build_item_card_set(session):
    rows = session.execute(text("SELECT number FROM items WHERE number IS NOT NULL")).fetchall()
    return {r[0] for r in rows}


def check_mode(rows):
    print("\n=== CHECK MODE: перші 5 рядків ===\n")
    for i, row in enumerate(rows[:5], start=3):
        print(f"--- Рядок {i} ---")
        for idx, label in COL_LABELS.items():
            val = col(row, idx)
            if val is not None and str(val).strip():
                print(f"  [{label}] = {repr(val)}")
    print("\n=== END CHECK ===\n")


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/import_movements.py <file.xlsx> [--check]")
        sys.exit(1)

    filepath = sys.argv[1]
    check = '--check' in sys.argv

    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        sys.exit(1)

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=3))
    print(f"Found {len(rows)} data rows")

    if check:
        check_mode(rows)
        return

    with Session(engine) as session:
        person_map    = build_person_map(session)
        item_card_set = build_item_card_set(session)
        print(f"Loaded {len(person_map)} persons, {len(item_card_set)} item cards from DB")

        admin_id = session.execute(
            text("SELECT id FROM users WHERE username = 'admin' LIMIT 1")
        ).scalar()

        # ── Pass 1: collect document keys ─────────────────────────────────
        doc_keys = {}
        movement_rows = []
        skipped = 0

        for i, row in enumerate(rows, start=3):
            item_name = clean(col(row, 1))
            if not item_name:
                skipped += 1
                continue

            doc_number  = clean(col(row, 13))
            doc_date    = parse_date(col(row, 12))
            op_type_raw = clean(col(row, 25))
            doc_type    = OP_TYPE_MAP.get((op_type_raw or '').lower(), op_type_raw or 'надходження')
            doc_key     = (doc_number, doc_date, doc_type)

            if doc_key not in doc_keys and (doc_number or doc_date):
                doc_keys[doc_key] = {
                    'from_unit': clean(col(row, 7)),
                    'to_unit':   clean(col(row, 8)),
                    'basis':     clean(col(row, 11)),
                    'service':   clean(col(row, 20)),
                }
            movement_rows.append((i, row, doc_key))

        print(f"Found {len(doc_keys)} unique documents in file")

        # ── Pass 2: get or create Document records ─────────────────────────
        doc_id_map   = {}
        created_docs = 0
        reused_docs  = 0

        for doc_key, hdr in doc_keys.items():
            doc_number, doc_date, doc_type = doc_key

            existing = session.execute(
                text("""SELECT id FROM documents
                        WHERE doc_number IS NOT DISTINCT FROM :num
                          AND doc_date   IS NOT DISTINCT FROM :date
                          AND doc_type   = :dtype
                        LIMIT 1"""),
                {'num': doc_number, 'date': doc_date, 'dtype': doc_type}
            ).scalar()

            if existing:
                doc_id_map[doc_key] = existing
                reused_docs += 1
            else:
                doc = Document(
                    doc_type   = doc_type,
                    doc_number = doc_number,
                    doc_date   = doc_date,
                    from_unit  = hdr['from_unit'],
                    to_unit    = hdr['to_unit'],
                    basis      = hdr['basis'],
                    service    = hdr['service'],
                    status     = 'signed',
                    signed_at  = datetime.now(timezone.utc),
                    created_by = admin_id,
                )
                session.add(doc)
                session.flush()
                doc_id_map[doc_key] = doc.id
                created_docs += 1

        print(f"Documents: {created_docs} created, {reused_docs} reused")

        # ── Pass 3: create Movement records ───────────────────────────────
        imported = 0
        unmatched_persons = set()

        for i, row, doc_key in movement_rows:
            mvo_from_name = clean(col(row, 9))
            mvo_to_name   = clean(col(row, 10))
            mvo_from_id = person_map.get(mvo_from_name.lower()) if mvo_from_name else None
            mvo_to_id   = person_map.get(mvo_to_name.lower())   if mvo_to_name   else None

            if mvo_from_name and not mvo_from_id:
                unmatched_persons.add(mvo_from_name)
            if mvo_to_name and not mvo_to_id:
                unmatched_persons.add(mvo_to_name)

            raw_card      = clean(col(row, 23))
            item_card_num = raw_card if raw_card in item_card_set else None

            m = Movement(
                document_id       = doc_id_map.get(doc_key),
                entry_date        = parse_date(col(row, 0)),
                item_name         = clean(col(row, 1)),
                item_card_num     = item_card_num,
                unit_of_measure   = clean(col(row, 3)),
                category          = clean(col(row, 4)),
                qty_in            = parse_decimal(col(row, 5)),
                qty_out           = parse_decimal(col(row, 6)),
                from_unit         = clean(col(row, 7)),
                to_unit           = clean(col(row, 8)),
                mvo_from_id       = mvo_from_id,
                mvo_to_id         = mvo_to_id,
                basis             = clean(col(row, 11)),
                doc_date          = parse_date(col(row, 12)),
                doc_number        = clean(col(row, 13)),
                serial_number     = clean(col(row, 15)),
                nomenclature_code = clean(col(row, 17)),
                price             = parse_decimal(col(row, 19)),
                service           = clean(col(row, 20)),
                doc_type          = clean(col(row, 24)),
                recipient_category= clean(col(row, 26)),
                created_by        = admin_id,
            )
            session.add(m)
            imported += 1

        session.commit()

        if unmatched_persons:
            print(f"\nWarning: МВО not found in persons ({len(unmatched_persons)}):")
            for name in sorted(unmatched_persons):
                print(f"  - {name}")

        print(f"\nDone: {imported} movements, {skipped} skipped, "
              f"{created_docs} docs created, {reused_docs} reused")


if __name__ == '__main__':
    main()
